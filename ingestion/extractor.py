"""
Extracts text and tables from PDF, DOCX, and Markdown files.

PDF:  pdfplumber for text + tables (tables converted to markdown grids)
DOCX: python-docx for paragraphs + tables
MD:   plain UTF-8 read

Returns a structured dict so the chunker and pipeline can work uniformly
regardless of source format.
"""
from __future__ import annotations
import io
import re
from pathlib import Path


def extract(filename: str, file_bytes: bytes) -> dict:
    """
    Extract content from file bytes.

    Returns:
        {
          "text":     str,            # full document text (tables inlined as markdown)
          "pages":    list[dict],     # per-page breakdown (page_num, text, tables)
          "metadata": dict,           # format-specific stats
          "warnings": list[str],
        }
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(file_bytes, filename)
    elif ext == ".docx":
        return _extract_docx(file_bytes, filename)
    elif ext in (".md", ".txt"):
        return _extract_markdown(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: .pdf, .docx, .md, .txt")


# ── PDF ───────────────────────────────────────────────────────────────────────

def _extract_pdf(file_bytes: bytes, filename: str) -> dict:
    import pdfplumber

    pages = []
    warnings = []

    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            total_pages = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                raw_text = page.extract_text() or ""
                tables_md = []
                for tbl in page.extract_tables():
                    md = _table_to_markdown(tbl)
                    if md:
                        tables_md.append(md)
                pages.append({
                    "page_num": i + 1,
                    "text": raw_text.strip(),
                    "tables": tables_md,
                })
    except Exception as exc:
        raise RuntimeError(f"Failed to read PDF '{filename}': {exc}") from exc

    if all(not p["text"] and not p["tables"] for p in pages):
        warnings.append(
            "PDF appears to be scanned / image-only — no text was extracted. "
            "OCR is not supported; please use a text-based PDF."
        )

    full_parts: list[str] = []
    for p in pages:
        if p["text"]:
            full_parts.append(f"[Page {p['page_num']}]\n{p['text']}")
        for tbl in p["tables"]:
            full_parts.append(tbl)

    return {
        "text": "\n\n".join(full_parts),
        "pages": pages,
        "metadata": {
            "total_pages": total_pages,
            "has_tables": any(p["tables"] for p in pages),
        },
        "warnings": warnings,
    }


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _extract_docx(file_bytes: bytes, filename: str) -> dict:
    from docx import Document

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise RuntimeError(f"Failed to open DOCX '{filename}': {exc}") from exc

    paragraphs: list[str] = []
    tables_md: list[str] = []

    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt:
            paragraphs.append(txt)

    for tbl in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
        md = _table_to_markdown(rows)
        if md:
            tables_md.append(md)

    full_text = "\n\n".join(paragraphs)
    if tables_md:
        full_text += "\n\n" + "\n\n".join(tables_md)

    return {
        "text": full_text,
        "pages": [{"page_num": 1, "text": full_text, "tables": tables_md}],
        "metadata": {
            "total_pages": None,
            "has_tables": bool(tables_md),
            "paragraph_count": len(paragraphs),
        },
        "warnings": [],
    }


# ── Markdown ──────────────────────────────────────────────────────────────────

def _extract_markdown(file_bytes: bytes, filename: str) -> dict:
    text = file_bytes.decode("utf-8", errors="replace")
    return {
        "text": text,
        "pages": [{"page_num": 1, "text": text, "tables": []}],
        "metadata": {
            "total_pages": 1,
            "has_tables": bool(re.search(r"^\|", text, re.MULTILINE)),
            "char_count": len(text),
        },
        "warnings": [],
    }


# ── Helpers ───────────────────────────────────────────────────────────────────

def _table_to_markdown(table: list[list]) -> str:
    """Convert a list-of-rows (list-of-cells) into a GitHub-flavoured markdown table."""
    if not table:
        return ""
    # Normalise: all cells → stripped strings
    table = [[str(c).strip() if c is not None else "" for c in row] for row in table]
    # Skip tables that are entirely empty
    if all(all(c == "" for c in row) for row in table):
        return ""

    max_cols = max(len(row) for row in table)
    table = [row + [""] * (max_cols - len(row)) for row in table]

    lines = [
        "| " + " | ".join(table[0]) + " |",
        "| " + " | ".join(["---"] * max_cols) + " |",
    ]
    for row in table[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)
